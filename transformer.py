# -*- coding: utf-8 -*-

"""
SCRIPT AVANZADO PARA CONVERTIR DOCUMENTOS WORD A JATS XML VALIDADO USANDO GEMINI CLI

Descripción:
Este script extrae contenido estructurado (texto, tablas e imágenes) de un
documento .docx. Genera placeholders específicos para tablas e imágenes,
guarda las imágenes en un directorio separado, y construye un prompt avanzado
para que Gemini genere JATS XML. Finalmente, valida el XML resultante contra el
DTD oficial de JATS Publishing 1.3.

Autor: Tu Nombre (adaptado de una solución de IA de Gemini)
Versión: 2.0
Fecha: 2025-10-09

Requisitos:
- Python 3.7+
- python-docx (pip install python-docx)
- lxml (pip install lxml)
- Gemini CLI instalada y autenticada.

Uso:
python word_a_jats_avanzado.py "ruta/a/articulo.docx" "ruta/de/salida.xml"
"""

import os
import re
import subprocess
import sys
import time
import urllib.request
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from lxml import etree

# --- Constantes ---
DTD_ZIP_URL = "https://public.nlm.nih.gov/projects/jats/publishing/1.3/JATS-Publishing-1-3-MathML3-DTD.zip"
WORKSPACE_ROOT = Path(__file__).resolve().parent
DTD_FILENAME = "JATS-journalpublishing1-3-mathml3.dtd"
DTD_LOCAL_FILE = WORKSPACE_ROOT / DTD_FILENAME
IMAGE_OUTPUT_DIR = WORKSPACE_ROOT / "imagenes_extraidas"


def extraer_contenido_estructurado(docx_path: str) -> str:
    """
    Extrae contenido de un .docx, incluyendo texto, tablas e imágenes.
    Las imágenes se guardan y se insertan placeholders en el texto.
    Las tablas se convierten a un formato de texto simple dentro de placeholders.

    Args:
        docx_path (str): Ruta al archivo .docx.

    Returns:
        str: Cadena de texto con el contenido y los placeholders.
             Retorna None si hay un error.
    """
    try:
        doc = Document(docx_path)
        content_parts = []
        image_counter = 1

        # Crear directorio para imágenes si no existe
        IMAGE_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

        # Usamos un iterador para procesar elementos del cuerpo (párrafos y tablas)
        for element in doc.element.body:
            if element.tag.endswith('p'):
                para = doc.paragraphs[int(element.xpath('count(preceding-sibling::w:p)'))]
                
                # Manejo de imágenes (inline shapes)
                # Esta es una forma de detectar imágenes. Puede necesitar ajustes.
                if 'graphicData' in para._p.xml:
                    for shape in para._p.xpath('.//pic:pic'):
                        # Extraer la relación de la imagen
                        rel_id = shape.xpath('.//a:blip/@r:embed')[0]
                        image_part = doc.part.related_parts[rel_id]
                        
                        # Guardar la imagen
                        image_filename = f"imagen_{image_counter}.{image_part.content_type.split('/')[-1]}"
                        image_path = IMAGE_OUTPUT_DIR / image_filename
                        with open(image_path, "wb") as f:
                            f.write(image_part.blob)
                        
                        # Añadir placeholder. El pie de foto se asume que es el texto del párrafo.
                        caption = para.text.strip()
                        content_parts.append(f'\n[IMAGEN-PLACEHOLDER file="{image_filename}" caption="{caption}"]\n')
                        image_counter += 1
                else:
                    content_parts.append(para.text)

            elif element.tag.endswith('tbl'):
                table = doc.tables[int(element.xpath('count(preceding-sibling::w:tbl)'))]
                table_content = []
                for row in table.rows:
                    row_content = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                    table_content.append(" | ".join(row_content))
                
                # Asumimos que el párrafo anterior a la tabla puede ser su título/caption
                caption = ""
                if content_parts and content_parts[-1].strip():
                    caption = content_parts[-1].strip()
                    content_parts[-1] = "" # Eliminarlo para no duplicar

                content_parts.append(f'\n[TABLA-PLACEHOLDER caption="{caption}" content="{"; ".join(table_content)}"]\n')

        return '\n'.join(content_parts)

    except FileNotFoundError:
        print(f"Error: El archivo '{docx_path}' no fue encontrado.", file=sys.stderr)
        return None
    except Exception as e:
        print(f"Error al leer o procesar el archivo Word: {e}", file=sys.stderr)
        return None


def construir_prompt_avanzado(texto_articulo: str) -> str:
    """
    Construye un prompt detallado para Gemini, incluyendo instrucciones
    para manejar los placeholders de imágenes y tablas.
    """
    prompt = f"""
    Actúa como un experto en etiquetado JATS (Journal Article Tag Suite) XML, versión 1.3, para SciELO.

    TAREA:
    Convierte el siguiente texto de un artículo científico, que contiene placeholders especiales para imágenes y tablas, a un archivo XML bien formado que cumpla con el estándar JATS.

    INSTRUCCIONES Y MEJORES PRÁCTICAS:
    1.  **Estructura General:** Raíz `<article>` con `xmlns:xlink="http://www.w3.org/1999/xlink"` y `xml:lang="es"`. Debe contener `<front>`, `<body>`, y `<back>`.
    2.  **Sección <front>:**
        *   Incluye `<article-title>` y, si está disponible en el texto, `<article-id pub-id-type="doi">`.
        *   En `<contrib-group>` declara cada autor con `<contrib contrib-type="author">`, colocando al inicio los superíndices (`<xref ref-type="aff" rid="affX">X</xref>`) que apunten a sus afiliaciones.
        *   Crea un `<aff>` por cada filiación, con `id` (`aff1`, `aff2`, ...) y un `<label>` numérico que coincida con los superíndices.
    *   Si el autor cuenta con correo electrónico, añádelo dentro del `<contrib>` como `<email>autor@dominio</email>`. Marca al responsable de correspondencia con `corresp="yes"`, coloca un `<xref ref-type="corresp" rid="cor1">*</xref>` al inicio, y registra el dato en `<author-notes><corresp id="cor1"><email>...</email></corresp></author-notes>`.
        *   Mantén `<abstract>` y `<kwd-group>` como en la especificación original.
    3.  **Sección <body>:**
        *   Usa `<sec>` para secciones con un `<title>`. Los párrafos deben ir en `<p>`.
        *   **Referencias en el cuerpo:** Identifica citas bibliográficas en formato numérico (Vancouver: `1`, `[1]`, `1-3`, etc.) y en formato autor-fecha (APA: `(Apellido, 2020)`, `(Apellido & Otro, 2019)`, etc.). Normaliza la cita reemplazando el texto original por elementos `<xref ref-type="bibr" rid="ID_DE_REFERENCIA">`. El contenido textual del `<xref>` debe reflejar el estilo original (por ejemplo, `[1]` o `(Apellido, 2020)`). Cada `<xref>` debe apuntar al `id` del `<ref>` correspondiente en la sección de bibliografía.
        *   **Manejo de Placeholders:**
            *   **Imágenes:** Si encuentras `[IMAGEN-PLACEHOLDER file="..." caption="..."]`, conviértelo a la siguiente estructura JATS:
                ```xml
                <fig id="f_ID_UNICO">
                  <label>Figura X</label>
                  <caption><p>Texto del pie de foto aquí</p></caption>
                  <graphic mimetype="image" xlink:href="NOMBRE_DEL_ARCHIVO_AQUI"/>
                </fig>
                ```
                Reemplaza los valores correspondientes. Genera un `id` único y una `label` secuencial.
            *   **Tablas:** Si encuentras `[TABLA-PLACEHOLDER caption="..." content="..."]`, donde el contenido es texto delimitado por `|` para columnas y `;` para filas, conviértelo a:
                ```xml
                <table-wrap id="t_ID_UNICO">
                  <label>Tabla X</label>
                  <caption><p>Título de la tabla aquí</p></caption>
                  <table>
                    <thead>
                      <tr>
                        <th>Encabezado 1</th>
                        <th>Encabezado 2</th>
                      </tr>
                    </thead>
                    <tbody>
                      <tr>
                        <td>Dato fila 1, col 1</td>
                        <td>Dato fila 1, col 2</td>
                      </tr>
                      <tr>
                        <td>Dato fila 2, col 1</td>
                        <td>Dato fila 2, col 2</td>
                      </tr>
                    </tbody>
                  </table>
                </table-wrap>
                ```
                Interpreta la primera fila del contenido como `<thead>` con `<th>` y las siguientes como `<tbody>` con `<td>`.
    4.  **Sección <back>:** Usa `<ref-list>` y `<ref>` para la bibliografía, desglosando con `<element-citation>`. Asegúrate de que cada `<ref>` tenga un `id` único (por ejemplo `ref1`, `ref2`, …) y que coincida con los atributos `rid` utilizados en los `<xref>` del cuerpo. Si detectas múltiples citas que apuntan a la misma referencia, reutiliza el mismo `id`.
    5.  **Reglas Finales:**
        *   El XML debe ser perfectamente bien formado.
        *   No incluyas explicaciones en la salida, solo el código XML.
        *   Codifica caracteres especiales (`&` como `&amp;`, etc.).

    TEXTO DEL ARTÍCULO A CONVERTIR:
    ---
    {texto_articulo}
    ---
    """
    return prompt


def _should_retry_gemini(stderr: str, stdout: str) -> bool:
    combined = f"{stderr}\n{stdout}".lower()
    retry_tokens = (
        "429",
        "ratelimit",
        "resource has been exhausted",
        "resource exhausted",
        "error when talking to gemini api",
    )
    return any(token in combined for token in retry_tokens)


def invocar_gemini_cli(prompt: str, max_attempts: int = 3, base_backoff: int = 20) -> dict:
    """Invoca la CLI de Gemini con reintentos controlados y devuelve metadatos de la ejecución."""

    command = ["gemini", "-p", prompt, "-o", "text"]
    last_result: dict = {'stdout': '', 'stderr': '', 'returncode': 1, 'elapsed': 0.0}

    for attempt in range(1, max_attempts + 1):
        try:
            start = time.time()
            result = subprocess.run(
                command,
                capture_output=True,
                text=True,
                check=False,
                encoding="utf-8",
                cwd=str(WORKSPACE_ROOT),
            )
            elapsed = time.time() - start
        except FileNotFoundError:
            msg = "Error: El comando 'gemini' no se encontró."
            print(msg, file=sys.stderr)
            return {'stdout': '', 'stderr': msg, 'returncode': 127, 'elapsed': 0.0}
        except Exception as exc:
            msg = f"Error inesperado al invocar a Gemini: {exc}"
            print(msg, file=sys.stderr)
            return {'stdout': '', 'stderr': msg, 'returncode': 1, 'elapsed': 0.0}

        stdout = result.stdout.strip() if result.stdout else ''
        stderr = result.stderr.strip() if result.stderr else ''
        last_result = {'stdout': stdout, 'stderr': stderr, 'returncode': result.returncode, 'elapsed': elapsed}

        if stdout and not _should_retry_gemini(stderr, stdout):
            return last_result

        if attempt >= max_attempts:
            break

        if _should_retry_gemini(stderr, stdout):
            wait_time = base_backoff * attempt
            print(f"Aviso: la respuesta de Gemini indicó un límite de recursos. Reintentando en {wait_time} segundos...", file=sys.stderr)
            time.sleep(wait_time)
            continue

        break

    return last_result


def extraer_resumen_tokens(gemini_meta: dict) -> dict:
    """
    Intenta extraer información de tokens desde stdout o stderr de la CLI.
    Devuelve un dict con campos posibles: total, prompt_tokens, completion_tokens,
    raw (texto fuente donde se buscó).
    """
    search_text = (gemini_meta.get('stderr', '') or '') + "\n" + (gemini_meta.get('stdout', '') or '')
    tokens_info = {'found': False, 'total': None, 'prompt_tokens': None, 'completion_tokens': None, 'raw': ''}

    # Intenta buscar un bloque JSON con token info
    try:
        m = re.search(r"\{[^\}]*token[^\}]*\}", search_text)
        if m:
            tokens_info['raw'] = m.group(0)
            # Intentar parsear números con regex
            total = re.search(r"total[_\s-]*tokens\D*(\d+)", m.group(0), re.I)
            pt = re.search(r"prompt[_\s-]*tokens\D*(\d+)", m.group(0), re.I)
            ct = re.search(r"completion[_\s-]*tokens\D*(\d+)", m.group(0), re.I)
            if total:
                tokens_info['total'] = int(total.group(1))
            if pt:
                tokens_info['prompt_tokens'] = int(pt.group(1))
            if ct:
                tokens_info['completion_tokens'] = int(ct.group(1))
            if tokens_info['total'] or tokens_info['prompt_tokens'] or tokens_info['completion_tokens']:
                tokens_info['found'] = True
                return tokens_info
    except Exception:
        pass

    # Buscar patrones simples en el texto (ej. "Total tokens: 1234")
    total_m = re.search(r"total\s*tokens\D*(\d+)", search_text, re.I)
    pt_m = re.search(r"prompt\s*tokens\D*(\d+)", search_text, re.I)
    ct_m = re.search(r"completion\s*tokens\D*(\d+)", search_text, re.I)
    if total_m or pt_m or ct_m:
        tokens_info['found'] = True
        tokens_info['raw'] = search_text
        if total_m:
            tokens_info['total'] = int(total_m.group(1))
        if pt_m:
            tokens_info['prompt_tokens'] = int(pt_m.group(1))
        if ct_m:
            tokens_info['completion_tokens'] = int(ct_m.group(1))
        return tokens_info

    # No se encontró info conocida
    tokens_info['raw'] = search_text
    return tokens_info


def validar_jats_xml(xml_content: str) -> tuple[bool, list]:
    """
    Valida una cadena de contenido XML contra el DTD de JATS.
    Descarga el DTD si no está disponible localmente.

    Args:
        xml_content (str): El contenido XML a validar.

    Returns:
        tuple[bool, list]: Un booleano indicando si la validación fue exitosa,
                           y una lista de errores si falló.
    """
    # Descargar y extraer DTD si es necesario
    # Verificar si el DTD existe en la raíz o en el subdirectorio estándar
    possible_dtd_locations = [
        WORKSPACE_ROOT / DTD_FILENAME,
        WORKSPACE_ROOT / "JATS-Publishing-1-3-MathML3-DTD" / DTD_FILENAME
    ]
    
    dtd_path = None
    for path in possible_dtd_locations:
        if path.exists():
            dtd_path = path
            break
            
    if not dtd_path:
        try:
            zip_path = WORKSPACE_ROOT / "jats_dtd.zip"
            if not zip_path.exists():
                print(f"Descargando el paquete DTD de JATS desde {DTD_ZIP_URL}...")
                urllib.request.urlretrieve(DTD_ZIP_URL, str(zip_path))
                print("Paquete ZIP descargado.")

            print("Extrayendo DTD...")
            with zipfile.ZipFile(zip_path, 'r') as z:
                z.extractall(WORKSPACE_ROOT)
            print("DTD extraído exitosamente.")
            
            # Re-verificar ubicación
            for path in possible_dtd_locations:
                if path.exists():
                    dtd_path = path
                    break
        except Exception as e:
            print(f"No se pudo descargar/extraer el DTD. No se puede validar. Error: {e}", file=sys.stderr)
            return False, [f"Fallo en la descarga/extracción del DTD: {e}"]

    if not dtd_path:
         return False, [f"No se encontró el archivo {DTD_FILENAME} después de la extracción."]

    print(f"Usando DTD en: {dtd_path}")

    try:
        # Asegurarse de que el contenido XML es bytes codificados en utf-8
        # Patch DOCTYPE to match the DTD we have
        xml_content_patched = re.sub(
            r'<!DOCTYPE.*?>',
            '<!DOCTYPE article PUBLIC "-//NLM//DTD JATS (Z39.96) Journal Publishing DTD with MathML3 v1.3 20210610//EN" "JATS-journalpublishing1-3-mathml3.dtd">',
            xml_content,
            flags=re.DOTALL
        )
        
        xml_bytes = xml_content_patched.encode('utf-8')

        # Parsear el XML y el DTD
        parser = etree.XMLParser(dtd_validation=False, no_network=False)
        xml_tree = etree.fromstring(xml_bytes, parser)
        dtd = etree.DTD(str(dtd_path))

        # Validar
        is_valid = dtd.validate(xml_tree)
        errors = dtd.error_log.filter_from_errors()
        return is_valid, errors

    except etree.XMLSyntaxError as e:
        return False, [f"Error de sintaxis XML (mal formado): {e}"]
    except Exception as e:
        return False, [f"Error inesperado durante la validación: {e}"]


def guardar_salida_xml(xml_content: str, output_path: str):
    # (Esta función es idéntica a la versión anterior)
    try:
        with open(output_path, 'w', encoding='utf-8') as f:
            # Añadir la declaración DOCTYPE para que los validadores sepan qué DTD usar
            doctype_declaration = f'<!DOCTYPE article PUBLIC "-//NLM//DTD JATS (Z39.96) Journal Publishing DTD v1.3 20210610//EN" "{DTD_FILENAME}">' 
            # Limpiar cualquier XML declaration que Gemini pudiera añadir
            if xml_content.startswith('<?xml'):
                xml_content = xml_content.split('?>', 1)[-1].strip()

            final_content = f'<?xml version="1.0" encoding="UTF-8"?>\n{doctype_declaration}\n{xml_content}'
            f.write(final_content)
        print(f"Archivo JATS XML guardado exitosamente en: {output_path}")
    except Exception as e:
        print(f"Error al guardar el archivo de salida: {e}", file=sys.stderr)

def main():
    """Función principal que orquesta el proceso."""
    if len(sys.argv) != 3:
        print("Uso: python word_a_jats_avanzado.py <entrada.docx> <salida.xml>", file=sys.stderr)
        sys.exit(1)

    input_path, output_path = sys.argv[1], sys.argv[2]

    # Preparar archivo de log paralelo: mismo nombre base, extensión .txt
    base_name = os.path.splitext(output_path)[0]
    log_path = f"{base_name}.txt"

    def log(msg: str, error: bool = False):
        timestamp = datetime.now().isoformat(sep=' ', timespec='seconds')
        line = f"[{timestamp}] {msg}\n"
        with open(log_path, 'a', encoding='utf-8') as lf:
            lf.write(line)
        if error:
            print(msg, file=sys.stderr)
        else:
            print(msg)

    log(f"Inicio de proceso para entrada: {input_path} salida: {output_path}")
    log("Paso 1/5: Procesando para extraer contenido estructurado...")
    structured_text = extraer_contenido_estructurado(input_path)
    if not structured_text:
        sys.exit(1)
    
    log(f"Paso 2/5: Construyendo prompt avanzado para Gemini...")
    prompt = construir_prompt_avanzado(structured_text)

    log("Paso 3/5: Invocando a la IA de Gemini para la conversión (esto puede tardar)...")
    gemini_result = invocar_gemini_cli(prompt)
    log(f"Gemini returncode: {gemini_result.get('returncode')} Elapsed: {gemini_result.get('elapsed'):.2f}s")
    if gemini_result.get('stderr'):
        log(f"Gemini stderr: {gemini_result.get('stderr')}", error=True)

    if gemini_result.get('returncode') != 0 or not gemini_result.get('stdout'):
        log("La conversión falló. No se generó XML.", error=True)
        # Guardar info de Gemini en el log y salir
        log("Salida completa de Gemini (stdout):")
        log(gemini_result.get('stdout', ''))
        log("Salida completa de Gemini (stderr):")
        log(gemini_result.get('stderr', ''))
        sys.exit(1)

    generated_xml = gemini_result.get('stdout')

    # Limpieza robusta: buscar bloque de código XML o contenido XML directo
    # 1. Intentar buscar bloque markdown ```xml ... ```
    match = re.search(r"```xml\s*(.*?)\s*```", generated_xml, re.DOTALL)
    if match:
        generated_xml = match.group(1)
    else:
        # 2. Si no hay bloque, intentar encontrar el inicio del XML (<?xml o <article)
        # Buscar el primer '<' que parezca inicio de XML
        match_xml = re.search(r"(<\?xml|<!DOCTYPE|<article).*", generated_xml, re.DOTALL)
        if match_xml:
            generated_xml = match_xml.group(0)
            # Limpiar posible cierre de markdown si quedó colgado
            if generated_xml.endswith("```"):
                generated_xml = generated_xml[:-3]
    
    generated_xml = generated_xml.strip()

    log("Paso 4/5: Validando el XML generado contra el DTD de JATS...")
    is_valid, errors = validar_jats_xml(generated_xml)
    if is_valid:
        log("¡Validación exitosa! El XML es conforme al DTD de JATS.")
    else:
        log("--- ADVERTENCIA: La validación del XML falló. ---", error=True)
        log("El archivo se guardará, pero necesita correcciones manuales.", error=True)
        log("Errores encontrados:")
        for error in errors:
            log(f"- {error}", error=True)

    log("Paso 5/5: Guardando el resultado...")
    guardar_salida_xml(generated_xml, output_path)

    # Extraer y registrar resumen de tokens desde la salida de Gemini
    tokens_summary = extraer_resumen_tokens(gemini_result)
    log("Resumen de tokens de Gemini:")
    if tokens_summary.get('found'):
        log(f"- total: {tokens_summary.get('total')}")
        log(f"- prompt_tokens: {tokens_summary.get('prompt_tokens')}")
        log(f"- completion_tokens: {tokens_summary.get('completion_tokens')}")
    else:
        log("- No se encontró información de tokens en la salida de Gemini.")
        # Adjuntar texto crudo para diagnóstico
        log("--- TEXTO CRUDO DE GEMINI ---")
        log(tokens_summary.get('raw', '')[:10000])

    log("Proceso completado.")

if __name__ == "__main__":
    main()